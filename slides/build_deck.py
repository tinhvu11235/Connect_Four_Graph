"""Builds slides/bieu_dien_do_thi.pptx: how each method turns a board into a graph.

Run:  python slides/build_deck.py [output.pptx]

Needs ``python-pptx``.  Side effects: overwrites the output file.  The file is a script:
helpers first, then one block per slide, then ``prs.save``.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("bieu_dien_do_thi.pptx")

INK = "1B2340"
MUTED = "4A5468"
BOARD = "1D3F8C"
BOARD_DARK = "142C66"
TINT = "EEF2FA"
RED = "E63946"
RED_LINE = "A4161A"
YEL = "F2B705"
YEL_LINE = "B98900"
TEAL = "0F8F7E"
PURPLE = "6C4AB6"
GREY = "B4BCCB"
EDGE = "C9D0DC"
FAINT = "DDE2EA"
WHITE = "FFFFFF"
HEAD = "Cambria"
BODY = "Calibri"

COLS, ROWS = 7, 6
DIRS4 = ((1, 0), (0, 1), (1, 1), (1, -1))
DIRS8 = DIRS4 + tuple((-dx, -dy) for dx, dy in DIRS4)

# The running example: red = side to move, yellow = opponent.
PIECES = {(2, 0): "R", (3, 0): "R", (3, 1): "Y"}


def heights(pieces):
    """Return the number of pieces in each column of ``pieces``.  Side effects: none."""

    return [sum(1 for (c, _r) in pieces if c == col) for col in range(COLS)]


LANDING = {(c, h) for c, h in enumerate(heights(PIECES)) if h < ROWS}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------- primitives
def rgb(hex_):
    """Return the python-pptx colour of a hex string like ``"1B2340"``.  Side effects: none."""

    return RGBColor.from_string(hex_)


def new_slide(bg=WHITE):
    """Append a blank slide with background ``bg`` to the deck and return it.  Side effects: adds a slide to ``prs``."""

    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(bg)
    return slide


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False):
    """Add an auto-shape at (x, y, w, h) inches with optional fill/outline; return it.  Side effects: adds a shape to ``slide``."""

    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return s


def rounded(slide, x, y, w, h, fill, radius=0.08, **kw):
    """Add a rounded rectangle (corner ``radius`` as a fraction); return it.  Side effects: adds a shape to ``slide``."""

    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, **kw)
    s.adjustments[0] = radius
    return s


def line(slide, x1, y1, x2, y2, color=EDGE, width=1.25, arrow=False, dash=False):
    """Add a straight connector, optionally dashed or with an arrow head; return it.  Side effects: adds a connector to ``slide``."""

    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return c


def text(slide, x, y, w, h, paragraphs, size=16, color=INK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, space_after=0, margin=0.0):
    """Add a text box; ``paragraphs`` is a str or a list of paragraphs, each a str or a list of (text, opts).

    ``opts`` may override font, size, bold, italic and color per run.  Side effects: adds a text box to ``slide``.
    """

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, opts in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = rgb(opts.get("color", color))
    return box


def disc(slide, cx, cy, d, fill, label="", label_color=WHITE, size=16, line_hex=None, line_w=1.0):
    """Add a circle of diameter ``d`` centred at (cx, cy), optionally labelled; return it.  Side effects: adds a shape to ``slide``."""

    s = shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line_hex, line_w=line_w)
    if label:
        tf = s.text_frame
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = BODY
        r.font.color.rgb = rgb(label_color)
    return s


def header(slide, kicker, title):
    """Add the small coloured kicker line and the slide title.  Side effects: adds two text boxes to ``slide``."""

    text(slide, 0.6, 0.38, 12.1, 0.35, kicker.upper(), size=13, color=TEAL, bold=True)
    text(slide, 0.6, 0.7, 12.1, 0.75, title, size=30, color=INK, font=HEAD, bold=True)


def card(slide, x, y, w, h, number, title, body, color=BOARD, body_size=15):
    """Add a numbered card (tinted box, number disc, title, body text).  Side effects: adds shapes to ``slide``."""

    rounded(slide, x, y, w, h, TINT)
    disc(slide, x + 0.45, y + 0.45, 0.5, color, str(number), size=16)
    text(slide, x + 0.9, y + 0.2, w - 1.1, 0.4, title, size=17, bold=True)
    text(slide, x + 0.9, y + 0.58, w - 1.1, h - 0.68, body, size=body_size, color=MUTED)


def notes(slide, body):
    """Set the speaker notes of ``slide``.  Side effects: overwrites the slide's notes."""

    slide.notes_slide.notes_text_frame.text = body


# ------------------------------------------------------------ board drawing
class Grid:
    """Maps cells (column, row; row 0 at the bottom) to slide inches."""

    def __init__(self, ox, oy, cell, cols=COLS, rows=ROWS):
        """Grid with top-left corner (ox, oy) inches and square cells of ``cell`` inches.  Side effects: none."""

        self.ox, self.oy, self.cell, self.cols, self.rows = ox, oy, cell, cols, rows

    def at(self, c, r):
        """Return the slide coordinates (inches) of the centre of cell (c, r).  Side effects: none."""

        return self.ox + (c + 0.5) * self.cell, self.oy + (self.rows - 1 - r + 0.5) * self.cell

    @property
    def w(self):
        """Grid width in inches.  Side effects: none."""

        return self.cols * self.cell

    @property
    def h(self):
        """Grid height in inches.  Side effects: none."""

        return self.rows * self.cell


def piece_node(slide, g, c, r, kind, d):
    """Draw one graph node of ``kind`` (R, Y, land, empty, absent) at cell (c, r) of grid ``g``.  Side effects: adds a shape to ``slide``."""

    x, y = g.at(c, r)
    if kind == "R":
        disc(slide, x, y, d, RED, line_hex=RED_LINE, line_w=1.25)
    elif kind == "Y":
        disc(slide, x, y, d, YEL, line_hex=YEL_LINE, line_w=1.25)
    elif kind == "land":
        disc(slide, x, y, d, WHITE, line_hex=TEAL, line_w=2.25).line.dash_style = MSO_LINE_DASH_STYLE.DASH
    elif kind == "empty":
        disc(slide, x, y, d, WHITE, line_hex=GREY, line_w=1.25)
    elif kind == "absent":
        disc(slide, x, y, d * 0.22, FAINT)


def cell_kind(c, r, pieces=PIECES, landing=LANDING):
    """Return the node kind of cell (c, r): the piece, ``"land"`` for a landing cell, else ``"empty"``.  Side effects: none."""

    if (c, r) in pieces:
        return pieces[(c, r)]
    return "land" if (c, r) in landing else "empty"


def physical_board(slide, g, pieces, landing=(), pad=0.14):
    """Draw a blue Connect 4 board with ``pieces`` and dashed ``landing`` cells on grid ``g``.  Side effects: adds shapes to ``slide``."""

    rounded(slide, g.ox - pad, g.oy - pad, g.w + 2 * pad, g.h + 2 * pad, BOARD, radius=0.05)
    d = g.cell * 0.78
    for c in range(g.cols):
        for r in range(g.rows):
            kind = pieces.get((c, r))
            x, y = g.at(c, r)
            if kind == "R":
                disc(slide, x, y, d, RED, line_hex=RED_LINE, line_w=1.25)
            elif kind == "Y":
                disc(slide, x, y, d, YEL, line_hex=YEL_LINE, line_w=1.25)
            else:
                s = disc(slide, x, y, d, WHITE, line_hex=TEAL if (c, r) in landing else BOARD_DARK,
                         line_w=2.5 if (c, r) in landing else 1.0)
                if (c, r) in landing:
                    s.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def graph_panel(slide, g, pad=0.25):
    """Draw the tinted background panel behind a graph on grid ``g``.  Side effects: adds a shape to ``slide``."""

    rounded(slide, g.ox - pad, g.oy - pad, g.w + 2 * pad, g.h + 2 * pad, TINT, radius=0.04)


def legend_item(slide, x, y, kind, label, size=13):
    """Draw one legend entry (node or edge symbol of ``kind`` + ``label``).  Side effects: adds shapes to ``slide``."""

    d = 0.24
    if kind == "R":
        disc(slide, x + d / 2, y + 0.15, d, RED, line_hex=RED_LINE)
    elif kind == "Y":
        disc(slide, x + d / 2, y + 0.15, d, YEL, line_hex=YEL_LINE)
    elif kind == "land":
        disc(slide, x + d / 2, y + 0.15, d, WHITE, line_hex=TEAL, line_w=2).line.dash_style = MSO_LINE_DASH_STYLE.DASH
    elif kind == "empty":
        disc(slide, x + d / 2, y + 0.15, d, WHITE, line_hex=GREY)
    elif kind == "absent":
        disc(slide, x + d / 2, y + 0.15, 0.08, GREY)
    elif kind in ("edge", "teal", "purple"):
        color = {"edge": GREY, "teal": TEAL, "purple": PURPLE}[kind]
        line(slide, x, y + 0.15, x + 0.3, y + 0.15, color=color, width=2.5 if kind != "edge" else 1.5)
    text(slide, x + 0.4, y, 3.0, 0.3, label, size=size, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)


# ===================================================================== slides
# 1 ── Title -----------------------------------------------------------------
s = new_slide(BOARD)
text(s, 0.8, 1.75, 6.8, 1.85, ["Biến bàn cờ", "thành đồ thị"], size=46, color=WHITE, font=HEAD, bold=True)
text(s, 0.8, 3.65, 6.2, 1.0, "Ba cách để mạng nơ-ron đồ thị (GNN) “nhìn” một ván Connect 4",
     size=20, color="CFDAF2")
text(s, 0.8, 6.5, 6.5, 0.4, "Connect4GNN · Phương pháp 1, Phương pháp 2 và AlphaGateau",
     size=13, color="9FB1DA")
g = Grid(8.1, 1.45, 0.66)
for c in range(COLS):
    for r in range(ROWS):
        for dx, dy in DIRS4:
            c2, r2 = c + dx, r + dy
            if 0 <= c2 < COLS and 0 <= r2 < ROWS:
                line(s, *g.at(c, r), *g.at(c2, r2), color="3E5DA8", width=1.0)
deco = {(2, 0): "R", (3, 0): "R", (3, 1): "Y", (4, 0): "Y", (4, 1): "R", (3, 2): "Y", (5, 0): "R"}
for c in range(COLS):
    for r in range(ROWS):
        x, y = g.at(c, r)
        kind = deco.get((c, r))
        if kind == "R":
            disc(s, x, y, 0.4, RED, line_hex=RED_LINE)
        elif kind == "Y":
            disc(s, x, y, 0.4, YEL, line_hex=YEL_LINE)
        else:
            disc(s, x, y, 0.4, BOARD_DARK, line_hex="5672B8")
notes(s, "Hôm nay mình trình bày cách biến một bàn cờ Connect 4 thành một đồ thị để mạng GNN đọc được. "
         "Có ba cách: hai phương pháp của nhóm và cách của bài báo AlphaGateau (vốn làm cho cờ vua).")

# 2 ── What is a graph -------------------------------------------------------
s = new_slide()
header(s, "Khởi động", "Đồ thị chỉ là những chấm và dây nối")
# diagram: one centre node listening to its neighbours
cx, cy = 3.3, 4.15
nbrs = [(1.4, 2.6), (5.2, 2.5), (1.2, 5.4), (5.4, 5.6), (3.3, 6.3)]
rounded(s, 0.6, 1.75, 5.5, 5.15, TINT, radius=0.04)
for nx, ny in nbrs:
    dxy = ((cx - nx) ** 2 + (cy - ny) ** 2) ** 0.5
    k = 0.42 / dxy
    line(s, nx, ny, cx - (cx - nx) * k, cy - (cy - ny) * k, color=TEAL, width=2.5, arrow=True)
line(s, 1.4, 2.6, 1.2, 5.4, color=GREY, width=1.5)
for nx, ny in nbrs:
    disc(s, nx, ny, 0.6, WHITE, line_hex=BOARD, line_w=2)
disc(s, cx, cy, 0.75, BOARD)
text(s, cx + 0.5, cy - 0.2, 1.6, 0.4, "nút", size=16, bold=True, color=BOARD)
text(s, 3.5, 5.15, 1.8, 0.4, "cạnh (dây)", size=16, bold=True, color=TEAL)
rows = [
    ("Nút (node)", "Một “nhân vật” giữ thông tin của riêng nó."),
    ("Cạnh (edge)", "Hai nút được phép “nói chuyện” với nhau."),
    ("GNN", "Mỗi vòng, mọi nút nghe các nút nối với mình rồi cập nhật hiểu biết."),
]
for i, (t, b) in enumerate(rows):
    y = 1.85 + i * 1.3
    disc(s, 7.15, y + 0.4, 0.6, [BOARD, TEAL, RED][i], str(i + 1), size=18)
    text(s, 7.7, y + 0.05, 5.0, 0.4, t, size=19, bold=True)
    text(s, 7.7, y + 0.47, 5.0, 0.7, b, size=16, color=MUTED)
text(s, 6.85, 5.85, 5.9, 0.9,
     "Giống một lớp học: mỗi bạn chỉ hỏi những bạn ngồi cạnh — sau vài vòng, cả lớp đều biết chuyện.",
     size=16, italic=True, color=BOARD)
notes(s, "Đồ thị gồm nút và cạnh. GNN làm việc theo vòng: ở mỗi vòng, mỗi nút thu thập tin từ các nút nối với nó "
         "(mũi tên xanh) rồi cập nhật. Vì vậy câu hỏi quan trọng nhất là: cái gì là nút, và nối cái gì với cái gì.")

# 3 ── Running example -------------------------------------------------------
s = new_slide()
header(s, "Ví dụ dùng xuyên suốt", "Cùng một thế cờ cho cả ba cách")
g = Grid(1.0, 1.95, 0.76)
physical_board(s, g, PIECES, LANDING)
for c in range(COLS):
    text(s, g.at(c, 0)[0] - 0.3, g.oy + g.h + 0.2, 0.6, 0.3, f"cột {c}", size=11, color=MUTED, align=PP_ALIGN.CENTER)
items = [
    ("R", "Quân đỏ: quân của bên sắp đi"),
    ("Y", "Quân vàng: quân của đối thủ"),
    ("land", "Vòng nét đứt: ô mà quân tiếp theo sẽ rơi vào (mỗi cột đúng 1 ô)"),
]
for i, (k, label) in enumerate(items):
    y = 2.1 + i * 0.75
    d = 0.42
    if k == "R":
        disc(s, 7.6, y + 0.3, d, RED, line_hex=RED_LINE)
    elif k == "Y":
        disc(s, 7.6, y + 0.3, d, YEL, line_hex=YEL_LINE)
    else:
        disc(s, 7.6, y + 0.3, d, WHITE, line_hex=TEAL, line_w=2.5).line.dash_style = MSO_LINE_DASH_STYLE.DASH
    text(s, 8.05, y, 4.7, 0.6, label, size=17, color=INK, anchor=MSO_ANCHOR.MIDDLE)
rounded(s, 7.3, 4.6, 5.45, 1.35, TINT)
text(s, 7.6, 4.78, 4.9, 1.1, [
    [("Lưu ý: ", {"bold": True, "color": BOARD}),
     ("Connect 4 chỉ có 7 nước đi — “thả quân vào cột 0…6”. Quân luôn rơi xuống ô thấp nhất còn trống.",
      {})]], size=16, color=INK)
notes(s, "Đây là thế cờ ví dụ: đỏ có 2 quân ở hàng dưới cùng, vàng có 1 quân. "
         "Các vòng nét đứt là 7 ô mà nước đi tiếp theo có thể rơi vào. Mọi slide sau đều vẽ lại thế cờ này.")

# 4 ── Three questions -------------------------------------------------------
s = new_slide()
header(s, "Cách đọc phần tiếp theo", "Mỗi cách đều trả lời ba câu hỏi")
qs = [
    ("Cái gì là nút?", "Mọi ô? Chỉ ô có quân? Hay còn thêm nút khác?"),
    ("Nút nào nối với nút nào?", "Chỉ hàng xóm sát bên, hay nối xa hơn? Dây có mang thông tin không?"),
    ("Nước đi được đọc ở đâu?", "Mạng phải chấm điểm 7 nước đi — điểm đó lấy từ nút hay từ dây?"),
]
for i, (q, b) in enumerate(qs):
    x = 0.6 + i * 4.15
    rounded(s, x, 1.85, 3.85, 3.3, TINT)
    disc(s, x + 0.75, 2.6, 0.9, [BOARD, TEAL, RED][i], str(i + 1), size=26)
    text(s, x + 0.35, 3.25, 3.2, 0.5, q, size=19, bold=True)
    text(s, x + 0.35, 3.8, 3.2, 1.3, b, size=15, color=MUTED)
order = [("Phương pháp 2", "Lưới cố định", BOARD), ("Phương pháp 1", "Đồ thị động", TEAL),
         ("AlphaGateau", "Cờ vua → Connect 4", PURPLE)]
for i, (a, b, col) in enumerate(order):
    x = 0.6 + i * 4.15
    rounded(s, x, 5.55, 3.85, 1.15, col)
    text(s, x + 0.3, 5.68, 3.3, 0.4, a, size=18, bold=True, color=WHITE)
    text(s, x + 0.3, 6.1, 3.3, 0.4, b, size=15, color=WHITE)
for i in range(2):
    x = 0.6 + (i + 1) * 4.15 - 0.3
    text(s, x - 0.02, 5.85, 0.34, 0.5, "›", size=30, color=GREY, bold=True, align=PP_ALIGN.CENTER)
notes(s, "Để so sánh dễ, với mỗi cách ta chỉ hỏi ba câu: nút là gì, nối gì với gì, và nước đi đọc ở đâu. "
         "Mình đi từ cách đơn giản nhất (phương pháp 2) trước, rồi tới phương pháp 1, cuối cùng là AlphaGateau.")

# 5 ── Method 2: grid --------------------------------------------------------
s = new_slide()
header(s, "Phương pháp 2 · Lưới cố định", "Mỗi ô là một nút, nối với 8 ô xung quanh")
g = Grid(0.95, 1.95, 0.78)
graph_panel(s, g)
focus = (5, 2)
for c in range(COLS):
    for r in range(ROWS):
        for dx, dy in DIRS4:
            c2, r2 = c + dx, r + dy
            if 0 <= c2 < COLS and 0 <= r2 < ROWS:
                line(s, *g.at(c, r), *g.at(c2, r2), color=EDGE, width=1.0)
for dx, dy in DIRS8:
    line(s, *g.at(*focus), *g.at(focus[0] + dx, focus[1] + dy), color=TEAL, width=3.5)
d = g.cell * 0.5
for c in range(COLS):
    for r in range(ROWS):
        piece_node(s, g, c, r, cell_kind(c, r), d)
x, y = g.at(*focus)
disc(s, x, y, d, WHITE, line_hex=TEAL, line_w=3.5)
card(s, 7.35, 1.75, 5.4, 1.4, 1, "Nút: cả 42 ô",
     "Ô trống cũng là nút — nó vẫn mang và truyền thông tin.")
card(s, 7.35, 3.35, 5.4, 1.65, 2, "Dây: 8 ô xung quanh",
     "Ví dụ ô viền xanh nối 8 hàng xóm. Đồ thị giống hệt nhau ở mọi thế cờ, chỉ thông tin trong nút đổi.",
     color=TEAL)
card(s, 7.35, 5.2, 5.4, 1.45, 3, "Nước đi: đọc ở nút",
     "Điểm của cột c lấy từ ô sẽ rơi vào của cột c (vòng nét đứt).", color=RED)
notes(s, "Phương pháp 2 là cách thẳng thắn nhất: 42 ô là 42 nút, kể cả ô trống. Mỗi ô nối với tối đa 8 ô "
         "xung quanh (ngang, dọc, chéo) — tổng 131 dây. Đồ thị không bao giờ đổi; thứ thay đổi theo ván là thông tin "
         "trong từng nút. Để chấm điểm cột c, mạng nhìn vào nút của ô sẽ rơi vào ở cột đó.")

# 6 ── Method 2: features and spreading -------------------------------------
s = new_slide()
header(s, "Phương pháp 2 · Lưới cố định", "Mỗi nút biết 7 điều, và tin lan ra từng vòng")
rounded(s, 0.6, 1.75, 5.9, 5.1, TINT, radius=0.04)
disc(s, 1.2, 2.3, 0.6, WHITE, line_hex=TEAL, line_w=2.5).line.dash_style = MSO_LINE_DASH_STYLE.DASH
text(s, 1.7, 2.05, 4.6, 0.5, "Ví dụ: ô cột 3, hàng 2", size=17, bold=True)
feats = [
    ("Có quân mình?", "không"),
    ("Có quân đối thủ?", "không"),
    ("Là ô trống?", "có"),
    ("Đặt được ngay?", "có"),
    ("Cao bao nhiêu?", "hàng 2 / 5"),
    ("Hàng chẵn hay lẻ?", "chẵn"),
    ("Cách cột giữa?", "0 (chính giữa)"),
]
for i, (q, v) in enumerate(feats):
    y = 2.85 + i * 0.55
    disc(s, 1.1, y + 0.17, 0.14, [BOARD, BOARD, BOARD, TEAL, MUTED, MUTED, MUTED][i])
    text(s, 1.35, y, 2.8, 0.36, q, size=16, color=INK)
    text(s, 4.1, y, 2.2, 0.36, v, size=16, color=BOARD, bold=True)
# spreading: which cells the example cell has heard from, on the example Connect 4 board
centre = (3, 2)
REACHED = "9FDCD2"
for k in range(3):
    gx = 7.0 + k * 2.0
    mg = Grid(gx, 2.05, 0.23)
    rounded(s, gx - 0.08, mg.oy - 0.08, mg.w + 0.16, mg.h + 0.16, BOARD, radius=0.06)
    for c in range(COLS):
        for r in range(ROWS):
            x, y = mg.at(c, r)
            dist = max(abs(c - centre[0]), abs(r - centre[1]))
            kind = PIECES.get((c, r))
            if dist == 0:
                disc(s, x, y, 0.18, TEAL, line_hex=WHITE, line_w=1.5)
            elif kind:
                reached = dist <= k + 1
                disc(s, x, y, 0.18, RED if kind == "R" else YEL,
                     line_hex=TEAL if reached else BOARD_DARK, line_w=2.25 if reached else 0.75)
            else:
                disc(s, x, y, 0.18, REACHED if dist <= k + 1 else WHITE, line_hex=BOARD_DARK, line_w=0.75)
    text(s, gx, 3.55, mg.w, 0.35, f"Sau vòng {k + 1}", size=15, bold=True, align=PP_ALIGN.CENTER)
disc(s, 7.2, 4.2, 0.22, TEAL, line_hex=BOARD_DARK, line_w=0.75)
text(s, 7.4, 4.05, 2.4, 0.3, "ô đang xét (cột 3, hàng 2)", size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
disc(s, 10.05, 4.2, 0.22, REACHED, line_hex=BOARD_DARK, line_w=0.75)
text(s, 10.25, 4.05, 2.5, 0.3, "ô đã gửi được tin tới nó", size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
rounded(s, 7.0, 4.65, 5.75, 2.2, TINT)
text(s, 7.3, 4.85, 5.2, 1.9, [
    [("Mỗi vòng tin chỉ đi được 1 ô. ", {"bold": True, "color": INK}),
     ("Sau 3 vòng, ô đang xét đã “nghe” được mọi ô cách nó tới 3 ô — đủ để nhận ra một hàng 4.", {})],
    [("Cách tính này giống hệt một mạng tích chập (CNN) 3×3, chỉ được kể lại theo ngôn ngữ đồ thị.",
      {"color": MUTED})],
], size=16, color=INK, space_after=8)
notes(s, "Mỗi nút mang 7 con số đơn giản: có quân mình không, có quân đối thủ không, trống không, đặt được ngay "
         "không, cao bao nhiêu, hàng chẵn/lẻ, cách cột giữa bao xa. Ở mỗi vòng, tin chỉ lan 1 ô, nên mạng cần vài "
         "vòng để nhìn ra một hàng 4. Thực chất phép tính này trùng với một lớp tích chập 3×3.")

# 7 ── Method 1: dynamic nodes ----------------------------------------------
s = new_slide()
header(s, "Phương pháp 1 · Đồ thị động", "Chỉ giữ những ô có ý nghĩa")
g = Grid(0.95, 1.95, 0.78)
graph_panel(s, g)
present = set(PIECES) | LANDING
highlight = ((3, 2), (5, 0))
for (c, r) in present:
    for dx, dy in DIRS4:
        for k in range(1, 4):
            c2, r2 = c + k * dx, r + k * dy
            if (c2, r2) in present:
                line(s, *g.at(c, r), *g.at(c2, r2), color=GREY, width=1.5)
                break
line(s, *g.at(*highlight[0]), *g.at(*highlight[1]), color=TEAL, width=4)
d = g.cell * 0.5
for c in range(COLS):
    for r in range(ROWS):
        piece_node(s, g, c, r, cell_kind(c, r) if (c, r) in present else "absent", d)
x, y = g.at(4, 1)
text(s, x + 0.3, y - 0.5, 1.6, 0.3, "bắc qua ô trống", size=12, color=TEAL, bold=True)
card(s, 7.35, 1.75, 5.4, 1.65, 1, "Nút: quân đã đặt + 7 ô sắp rơi vào",
     "Ở ví dụ: 3 quân + 7 ô = 10 nút. 32 ô còn lại không có mặt (chấm mờ).")
card(s, 7.35, 3.6, 5.4, 1.4, 2, "Đồ thị lớn dần theo ván",
     "Bàn trống có 7 nút, bàn đầy có 42 nút.", color=TEAL)
card(s, 7.35, 5.2, 5.4, 1.45, 3, "Nước đi: vẫn đọc ở nút",
     "Ở 7 ô sắp rơi vào — vì thế các ô này luôn được giữ lại.", color=RED)
notes(s, "Phương pháp 1 bỏ đi những ô trống không quan trọng. Chỉ còn quân đã đặt và 7 ô sắp rơi vào — ở ví dụ là 10 nút. "
         "Ván càng dài đồ thị càng lớn. 7 ô sắp rơi vào bắt buộc phải giữ vì mạng chấm điểm nước đi tại đó. "
         "Dây nối các nút cách nhau tối đa 3 ô trên cùng một đường thẳng, kể cả khi ở giữa là ô không có mặt "
         "(dây xanh). Hình chỉ vẽ dây gần nhất mỗi hướng; dây xa hơn nằm chồng lên cùng đường.")

# 8 ── Method 1: long edges ------------------------------------------------------
s = new_slide()
header(s, "Phương pháp 1 · Đồ thị động", "Dây bắc qua ô trống và ghi lại “giữa có mấy ô trống”")
mini = {(0, 0): "Y", (1, 0): "Y", (3, 0): "Y", (0, 1): "R", (1, 1): "R", (3, 1): "R"}
g = Grid(1.05, 2.45, 1.25, cols=4, rows=2)
graph_panel(s, g, pad=0.35)
line(s, *g.at(0, 1), *g.at(1, 1), color=GREY, width=2)
line(s, *g.at(0, 0), *g.at(1, 0), color=GREY, width=2)
line(s, *g.at(1, 0), *g.at(2, 0), color=GREY, width=2)
line(s, *g.at(2, 0), *g.at(3, 0), color=GREY, width=2)
for c in (0, 1, 3):
    line(s, *g.at(c, 0), *g.at(c, 1), color=GREY, width=2)
line(s, *g.at(1, 1), *g.at(3, 1), color=TEAL, width=5)
d = g.cell * 0.5
for (c, r), k in mini.items():
    piece_node(s, g, c, r, k, d)
piece_node(s, g, 2, 0, "land", d)
x, y = g.at(2, 1)
shape(s, MSO_SHAPE.OVAL, x - d / 2, y - d / 2, d, d, fill=None, line=GREY, line_w=1.0, dash=True)
text(s, g.ox - 0.3, 1.6, 5.2, 0.45, "Mẫu “X X _ X”: đỏ chỉ thiếu 1 quân để thắng", size=16, bold=True, color=RED)
text(s, x - 0.9, y - 0.95, 1.8, 0.35, "1 ô trống ở giữa", size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
text(s, x - 1.2, g.oy + g.h + 0.5, 2.4, 0.55, "ô nét đứt xám: không phải nút", size=13, color=MUTED, align=PP_ALIGN.CENTER)
card(s, 6.9, 1.75, 5.85, 1.55, 1, "Vấn đề khi bỏ ô trống",
     "Ô giữa hai quân đỏ không có mặt, nên hai quân này không còn “nói chuyện” với nhau qua hàng xóm.")
card(s, 6.9, 3.5, 5.85, 1.45, 2, "Cách giải: dây nối xa tối đa 3 ô",
     "Hai nút cùng hàng, cột hay đường chéo, cách ≤ 3 ô, thì được nối thẳng.", color=TEAL)
card(s, 6.9, 5.15, 5.85, 1.5, 3, "Dây mang theo 1 con số",
     "Số ô trống nằm giữa hai đầu dây. Hướng và khoảng cách (1, 2, 3) được học riêng.", color=RED)
notes(s, "Bỏ ô trống thì có cái giá: hai quân đỏ cách nhau một ô trống không còn là hàng xóm. Phương pháp 1 bù lại bằng "
         "dây nối xa tối đa 3 ô trên cùng đường thẳng, và mỗi dây ghi số ô trống ở giữa. Nhờ vậy mạng vẫn nhận ra "
         "mẫu X X _ X — một mối đe dọa thắng.")

# 9 ── AlphaGateau: chess ------------------------------------------------------
s = new_slide()
header(s, "AlphaGateau · Cờ vua", "Mỗi dây nối là một nước đi")
g = Grid(1.3, 1.85, 0.86, cols=5, rows=5)
for c in range(5):
    for r in range(5):
        shape(s, MSO_SHAPE.RECTANGLE, g.ox + c * g.cell, g.oy + (4 - r) * g.cell, g.cell, g.cell,
              fill="C9D3E8" if (c + r) % 2 == 0 else "F3F5FA")
centre = (2, 2)
cx, cy = g.at(*centre)
queen = [(centre[0] + k * dx, centre[1] + k * dy) for dx, dy in DIRS8 for k in (1, 2)]
knight = [(centre[0] + a, centre[1] + b) for a, b in
          ((1, 2), (2, 1), (2, -1), (1, -2), (-1, -2), (-2, -1), (-2, 1), (-1, 2))]


def arrow_to(slide, c, r, color, width, dash=False):
    """Draw an arrow from the knight's square to cell (c, r), stopping short of the target.  Side effects: adds a connector to ``slide``."""

    tx, ty = g.at(c, r)
    dist = ((tx - cx) ** 2 + (ty - cy) ** 2) ** 0.5
    k = 0.2 / dist
    line(slide, cx, cy, tx - (tx - cx) * k, ty - (ty - cy) * k, color=color, width=width, arrow=True, dash=dash)


for c, r in queen:
    if max(abs(c - 2), abs(r - 2)) == 2:
        arrow_to(s, c, r, MUTED, 1.5, dash=True)
for c, r in knight:
    arrow_to(s, c, r, PURPLE, 3)
disc(s, cx, cy, 0.62, WHITE, line_hex=INK, line_w=1.5)
text(s, cx - 0.4, cy - 0.33, 0.8, 0.62, "♞", size=30, font="Segoe UI Symbol", color=INK, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
legend_item(s, 1.0, 6.55, "purple", "nước đi hợp lệ lúc này (mã)")
line(s, 4.15, 6.7, 4.45, 6.7, color=MUTED, width=1.5, dash=True)
text(s, 4.55, 6.55, 2.9, 0.3, "có thể có, nhưng không hợp lệ", size=13, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
card(s, 7.35, 1.7, 5.4, 1.2, 1, "Nút: mọi ô", "64 ô (bản nhỏ 5×5: 25 ô), nhớ cả 8 thế cờ gần nhất.",
     color=PURPLE)
card(s, 7.35, 3.05, 5.4, 1.45, 2, "Dây: mọi nước đi có thể có",
     "Vẽ sẵn một lần, không đổi theo ván: 1.858 dây với bàn 8×8.", color=PURPLE)
card(s, 7.35, 4.65, 5.4, 1.2, 3, "Dây ghi: hợp lệ? đi bao xa? quân nào?",
     "Chỉ thông tin “hợp lệ” thay đổi theo thế cờ.", color=PURPLE, body_size=14)
card(s, 7.35, 6.0, 5.4, 1.1, 4, "Nước đi: đọc ngay trên dây", "", color=PURPLE)
notes(s, "AlphaGateau làm cho cờ vua, và ý tưởng chính rất khác: dây nối chính là nước đi. Từ một ô, vẽ sẵn dây tới mọi ô "
         "mà một quân bất kỳ có thể đi tới (đường xe, tượng, hậu, bước mã…). Dây được vẽ một lần cho mọi thế cờ; "
         "thứ thay đổi là thông tin trên dây, ví dụ nước này lúc này có hợp lệ không. Vì nước đi là dây, "
         "điểm của nước đi cũng được đọc trên dây.")

# 10 ── Why it suits chess ---------------------------------------------------
s = new_slide()
header(s, "AlphaGateau · Cờ vua", "Vì sao cách này hợp với cờ vua")
stats = [
    ("1 vòng", "đã nhìn xa cả bàn", "Xe ở a1 nghe được a8 ngay, vì có dây nối thẳng — không phải chờ tin lan từng ô."),
    ("5×5 = 8×8", "cùng một mạng", "Dây chỉ ghi “đi bao xa, theo hướng nào”, không gắn với ô cụ thể, nên đổi cỡ bàn vẫn chạy."),
    ("Dây = nước", "chấm điểm tự nhiên", "Nước đi trong cờ vua là “từ ô A tới ô B” — đúng bằng một dây."),
]
for i, (big, small, body) in enumerate(stats):
    x = 0.6 + i * 4.15
    rounded(s, x, 1.85, 3.85, 3.4, TINT)
    text(s, x + 0.35, 2.15, 3.2, 1.0, big, size=40, bold=True, color=PURPLE, font=HEAD)
    text(s, x + 0.35, 3.15, 3.2, 0.45, small, size=18, bold=True, color=INK)
    text(s, x + 0.35, 3.75, 3.2, 1.4, body, size=16, color=MUTED)
rounded(s, 0.6, 5.65, 12.15, 1.0, BOARD)
text(s, 0.95, 5.65, 11.5, 1.0, "Cả ba lợi thế đều dựa vào một điều: cờ vua có rất nhiều kiểu nước đi “từ ô này sang ô kia”. "
     "Connect 4 thì không.", size=17, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Ba lợi ích: chỉ một vòng đã nhìn xa cả bàn vì dây nối thẳng; cùng một mạng chạy được nhiều cỡ bàn; "
         "và việc chấm điểm nước đi rất tự nhiên vì nước đi chính là dây.")

# 11 ── The problem in Connect 4 ---------------------------------------------
s = new_slide()
header(s, "AlphaGateau → Connect 4", "Vấn đề: thả quân thì không có ô xuất phát")
# chess panel
rounded(s, 0.6, 1.75, 5.9, 3.65, TINT, radius=0.04)
text(s, 0.9, 1.9, 5.3, 0.4, "Cờ vua: đi từ ô A tới ô B", size=18, bold=True, color=PURPLE)
cg = Grid(2.0, 2.5, 0.75, cols=4, rows=3)
for c in range(4):
    for r in range(3):
        shape(s, MSO_SHAPE.RECTANGLE, cg.ox + c * cg.cell, cg.oy + (2 - r) * cg.cell, cg.cell, cg.cell,
              fill="C9D3E8" if (c + r) % 2 == 0 else WHITE)
ax, ay = cg.at(0, 0)
bx, by = cg.at(3, 2)
line(s, ax, ay, bx - 0.2, by + 0.13, color=PURPLE, width=4, arrow=True)
disc(s, ax, ay, 0.42, WHITE, "A", label_color=INK, line_hex=INK)
disc(s, bx, by, 0.42, WHITE, "B", label_color=INK, line_hex=INK)
# connect 4 panel
rounded(s, 6.85, 1.75, 5.9, 3.65, TINT, radius=0.04)
text(s, 7.15, 1.9, 5.3, 0.4, "Connect 4: thả vào cột — từ đâu?", size=18, bold=True, color=RED)
kg = Grid(8.7, 3.3, 0.6, cols=4, rows=3)
physical_board(s, kg, {(1, 0): "Y", (2, 0): "R", (2, 1): "Y"}, {(1, 1)}, pad=0.1)
tx, ty = kg.at(1, 2)
disc(s, tx, ty - 0.95, 0.45, RED, line_hex=RED_LINE)
line(s, tx, ty - 0.68, tx, ty - 0.08, color=RED, width=3, arrow=True)
text(s, tx + 0.35, ty - 1.15, 1.0, 0.45, "?", size=28, bold=True, color=RED)
rounded(s, 0.6, 5.7, 12.15, 1.15, WHITE, line=GREY, line_w=1.0)
text(s, 0.95, 5.85, 11.5, 0.9, [
    [("Bê nguyên “dây = nước đi” sang Connect 4: ", {"bold": True, "color": INK}),
     ("không có dây nào nối các ô với nhau, nên các ô không thể “nói chuyện” — mạng không nhìn ra hàng 4.", {})]],
     size=17, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Khi đem sang Connect 4 thì vấp ngay một vấn đề: nước đi ở đây là thả quân vào cột, không có ô xuất phát. "
         "Nếu chỉ vẽ dây theo nước đi thì các ô trên bàn không được nối với nhau, mạng sẽ không thấy được hàng 4.")

# 12 ── Fix: column nodes + two kinds of edges ----------------------------------
s = new_slide()
header(s, "AlphaGateau → Connect 4", "Cách sửa: thêm 7 nút “cột” và hai loại dây")
g = Grid(1.05, 2.45, 0.7)
graph_panel(s, g, pad=0.2)
col_y = 1.85
focus = (4, 2)
for dx, dy in DIRS8:
    far = None
    for k in range(1, 4):
        c2, r2 = focus[0] + k * dx, focus[1] + k * dy
        if 0 <= c2 < COLS and 0 <= r2 < ROWS:
            far = (c2, r2)
    if far:
        line(s, *g.at(*focus), *g.at(*far), color=PURPLE, width=2.5)
for c in range(COLS):
    x = g.at(c, 0)[0]
    h = heights(PIECES)[c]
    lx, ly = g.at(c, h)
    line(s, x, col_y + 0.2, lx, ly - g.cell * 0.27, color=TEAL, width=2.5, arrow=True)
d = g.cell * 0.5
for c in range(COLS):
    for r in range(ROWS):
        kind = cell_kind(c, r)
        piece_node(s, g, c, r, kind, d)
x, y = g.at(*focus)
disc(s, x, y, d, WHITE, line_hex=PURPLE, line_w=3)
for c in range(COLS):
    x = g.at(c, 0)[0]
    rounded(s, x - 0.22, col_y - 0.22, 0.44, 0.44, TEAL, radius=0.2)
    text(s, x - 0.22, col_y - 0.22, 0.44, 0.44, str(c), size=14, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
legend_item(s, 1.0, 6.85, "teal", "dây “thả” đang hợp lệ")
legend_item(s, 4.0, 6.85, "purple", "dây “đường thắng”")
card(s, 7.35, 1.7, 5.4, 1.2, 1, "Nút: 42 ô + 7 nút cột", "Nút cột đóng vai “ô xuất phát” của nước thả.",
     color=PURPLE)
card(s, 7.35, 3.05, 5.4, 1.45, 2, "Dây “thả”: cột → ô",
     "42 dây cố định; dây tới ô sắp rơi vào được đánh dấu “hợp lệ”.", color=TEAL)
card(s, 7.35, 4.65, 5.4, 1.45, 3, "Dây “đường thắng”: ô ↔ ô",
     "Nối hai ô cách ≤ 3 trên một đường thẳng — có thể cùng nằm trong một hàng 4.", color=PURPLE)
card(s, 7.35, 6.25, 5.4, 0.9, 4, "Nước đi: đọc trên dây “thả” hợp lệ", "", color=RED)
notes(s, "Cách dịch hợp lý: thêm 7 nút đại diện cho 7 cột. Dây thả nối nút cột tới các ô trong cột; dây đi tới ô sắp rơi vào "
         "được đánh dấu hợp lệ — giống cờ vua chỉ đổi cờ hợp lệ trên dây. Để các ô nói chuyện với nhau, thêm dây đường "
         "thắng nối các ô cách nhau tối đa 3 trên một đường thẳng (vẽ cho ô viền tím). Điểm của nước đi đọc trên dây thả hợp lệ.")

# 13 ── Gains and losses -----------------------------------------------------------
s = new_slide()
header(s, "AlphaGateau → Connect 4", "Đem sang thì được gì, mất gì")
pros = [
    ("Dây có “trí nhớ” riêng", "Dây được cập nhật mỗi vòng, ví dụ học “đoạn này còn mở hay đã bị chặn”."),
    ("Tự chọn nghe ai nhiều hơn", "Cơ chế attention cho mỗi nút chú ý tới dây quan trọng."),
    ("Đổi cỡ bàn không cần sửa", "Dây chỉ ghi hướng và khoảng cách."),
]
cons = [
    ("Ý tưởng chính gần như vô dụng", "Chỉ có 7 nước, mỗi nước tới đúng 1 ô — đọc ở dây hay ở ô cũng như nhau."),
    ("Chậm hơn khi tự chơi", "Phép tính kiểu đồ thị thưa chậm hơn tích chập trên bàn nhỏ 7×6."),
    ("Dây “đường thắng” ≈ dây nhìn xa của PP1", "Khác ở chỗ vẫn giữ đủ 42 ô, không bỏ ô trống."),
]
for col, (items, title, color, sign) in enumerate(((pros, "Được", TEAL, "+"), (cons, "Mất", RED, "−"))):
    x = 0.6 + col * 6.2
    disc(s, x + 0.3, 1.95, 0.5, color, sign, size=22)
    text(s, x + 0.7, 1.72, 4.0, 0.45, title, size=22, bold=True, color=color)
    for i, (t, b) in enumerate(items):
        y = 2.5 + i * 1.3
        rounded(s, x, y, 5.95, 1.15, TINT)
        text(s, x + 0.3, y + 0.12, 5.4, 0.4, t, size=17, bold=True)
        text(s, x + 0.3, y + 0.52, 5.4, 0.6, b, size=15, color=MUTED)
rounded(s, 0.6, 6.4, 12.15, 0.7, BOARD)
text(s, 0.9, 6.4, 11.6, 0.7, "Tóm lại: bản Connect 4 của AlphaGateau nằm giữa PP2 và PP1 — "
     "điểm mới thật sự là dây có thông tin riêng và attention.", size=16, color=WHITE, bold=True,
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Lợi: dây có thông tin riêng được cập nhật mỗi vòng, có attention, và chạy được nhiều cỡ bàn. "
         "Hại: ý tưởng dây = nước đi gần như không giúp gì vì Connect 4 chỉ có 7 nước, mỗi nước tới đúng một ô; "
         "và chậm hơn khi self-play. Dây đường thắng thực chất gần giống dây nhìn xa của phương pháp 1.")

# 14 ── Comparison table ---------------------------------------------------------
s = new_slide()
header(s, "Tổng kết", "Ba cách đặt cạnh nhau")
data = [
    ["", "Phương pháp 2\nLưới cố định", "Phương pháp 1\nĐồ thị động", "AlphaGateau"],
    ["Nút là gì?", "Mọi ô (42)", "Quân đã đặt + 7 ô sắp rơi vào", "Mọi ô (cờ vua: 64)"],
    ["Nối gì với gì?", "8 ô xung quanh", "Ô cùng đường thẳng, cách ≤ 3", "Mọi nước đi có thể có"],
    ["Đồ thị đổi theo ván?", "Không", "Có", "Không (chỉ “hợp lệ” đổi)"],
    ["Dây mang thông tin?", "Không", "Số ô trống ở giữa", "Nhiều: hợp lệ, hướng, loại quân"],
    ["Nước đi đọc ở đâu?", "Ô sắp rơi vào", "Ô sắp rơi vào", "Trên dây"],
]
tbl = s.shapes.add_table(len(data), 4, Inches(0.6), Inches(1.75), Inches(12.15), Inches(5.0)).table
widths = [2.75, 3.0, 3.2, 3.2]
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
head_colors = [WHITE, BOARD, TEAL, PURPLE]
for i, row in enumerate(data):
    tbl.rows[i].height = Inches(1.0 if i == 0 else 0.8)
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        if i == 0:
            cell.fill.fore_color.rgb = rgb(head_colors[j])
        else:
            cell.fill.fore_color.rgb = rgb(TINT if i % 2 else WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.15)
        tf = cell.text_frame
        tf.word_wrap = True
        lines_ = val.split("\n")
        for k, ln_text in enumerate(lines_):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = ln_text
            r.font.name = BODY
            r.font.size = Pt(17 if i == 0 and k == 0 else 15)
            r.font.bold = i == 0 and k == 0 or j == 0
            r.font.color.rgb = rgb(WHITE if i == 0 else INK)
notes(s, "Bảng này gom lại ba câu hỏi ban đầu cho cả ba cách.")

# 15 ── Closing ---------------------------------------------------------------------
s = new_slide(BOARD)
text(s, 0.8, 0.9, 11.5, 0.9, "Ba cách, ba câu tóm tắt", size=40, color=WHITE, font=HEAD, bold=True)
summary = [
    ("Phương pháp 2 · Lưới", "Mọi ô đều có mặt, mỗi ô chỉ nói chuyện với hàng xóm sát bên.", WHITE),
    ("Phương pháp 1 · Động", "Chỉ giữ ô quan trọng, nối xa tối đa 3 ô để không bỏ sót mối đe dọa.", "7FE0D0"),
    ("AlphaGateau", "Dây nối chính là nước đi — rất hợp cờ vua, ít lợi thế với Connect 4.", "C7B5F0"),
]
for i, (t, b, col) in enumerate(summary):
    y = 2.35 + i * 1.45
    disc(s, 1.2, y + 0.45, 0.7, BOARD_DARK, str(i + 1), label_color=col, size=22, line_hex=col, line_w=2)
    text(s, 1.85, y + 0.05, 10.5, 0.45, t, size=21, bold=True, color=col)
    text(s, 1.85, y + 0.52, 10.5, 0.5, b, size=18, color="DCE4F5")
notes(s, "Kết lại bằng ba câu tóm tắt. Cảm ơn mọi người đã lắng nghe.")

prs.save(OUT)
print(f"wrote {OUT}")
